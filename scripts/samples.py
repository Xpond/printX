"""Generate small, repeatable fixtures for manual VM testing."""
import math
import random
from pathlib import Path


def generate(directory):
    from PIL import Image, ImageDraw, ImageOps
    import pymupdf
    from docx import Document
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    photo = Image.new('RGB', (1600, 1000))
    pixels = photo.load()
    for y in range(photo.height):
        for x in range(photo.width):
            hill = 570 + 130 * math.sin(x / 310)
            pixels[x, y] = (int(80 + y / 12), int(130 + y / 15), 205) if y < hill else (42, 105 + y // 25, 65)
    ImageDraw.Draw(photo).ellipse((1080, 120, 1220, 260), fill='#ffe3a5')
    photo.save(directory / '01 landscape.jpg', quality=95, dpi=(300, 300))
    logo = Image.new('RGBA', (300, 300), (0, 0, 0, 0))
    draw = ImageDraw.Draw(logo)
    draw.rectangle((30, 40, 180, 240), fill='#009cb4')
    draw.ellipse((100, 70, 270, 240), fill='#bd2368')
    logo.save(directory / '02 transparent logo.png')
    photo.resize((400, 250)).save(directory / '03 two-page.tiff', save_all=True,
        append_images=[Image.new('RGB', (250, 400), '#e3ad42')])
    for locked in (False, True):
        with pymupdf.open() as doc:
            for i in range(3):
                page = doc.new_page(width=595 if i < 2 else 420, height=842 if i < 2 else 595)
                page.insert_text((50, 80), f'PrintShop sample - page {i + 1}', fontsize=24)
                page.draw_rect((50, 120, 350, 300), color=(0, .5, .6), width=3)
            kwargs = dict(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw='printshop', owner_pw='owner') if locked else {}
            doc.save(directory / ('05 password-printshop.pdf' if locked else '04 three pages.pdf'), **kwargs)
    (directory / '06 broken.pdf').write_bytes(b'This is intentionally not a valid PDF.')
    doc = Document()
    doc.add_heading('PrintShop sample document', 0)
    doc.add_paragraph('Office conversion arrives in phase 4. This file tests the current unsupported-file message.')
    doc.save(directory / '07 office sample.docx')
    modules, pick = Image.new('1', (29, 29), 1), random.Random(8)
    for y in range(29):
        for x in range(29):
            modules.putpixel((x, y), pick.random() < .5)
    draw = ImageDraw.Draw(modules)
    for x, y in ((0, 0), (22, 0), (0, 22)):  # Corner squares like a QR code; it does not scan.
        draw.rectangle((x, y, x + 6, y + 6), fill=0)
        draw.rectangle((x + 1, y + 1, x + 5, y + 5), fill=1)
        draw.rectangle((x + 2, y + 2, x + 4, y + 4), fill=0)
    ImageOps.expand(modules, 4, fill=1).resize((148, 148), Image.NEAREST).save(directory / '08 qr-style code.png')


if __name__ == '__main__':
    import sys
    generate(sys.argv[1] if len(sys.argv) > 1 else 'samples')
